"""
REPORT EXPORT
==============
Builds downloadable Excel (multi-sheet) and PDF (summary one-pager)
reports for whatever Partner/RM/Cluster/Region + period is currently
selected in the dashboard. Returns raw bytes so app.py can hand them
straight to st.download_button — nothing is written to disk.
"""

import io
from datetime import date

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from formatting import format_inr, format_count, format_pct

NAVY = colors.HexColor("#0B3D66")
GOLD = colors.HexColor("#C89B3C")


def build_excel_report(header_name: str, period_label: str, kpis: dict, target_summary: pd.DataFrame,
                        product_df: pd.DataFrame, clients: pd.DataFrame, opportunities: dict,
                        reviews: pd.DataFrame) -> bytes:
    """Multi-sheet Excel workbook: KPI summary, Target vs Actual, Product
    Performance, Client Opportunities, Review History."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pd.DataFrame([kpis]).to_excel(writer, sheet_name="KPI Summary", index=False)
        target_summary.to_excel(writer, sheet_name="Target vs Actual", index=False)
        product_df.to_excel(writer, sheet_name="Product Performance", index=False)
        for name, df in opportunities.items():
            sheet = f"Opp - {name}"[:31]
            df.to_excel(writer, sheet_name=sheet, index=False)
        reviews.to_excel(writer, sheet_name="Review History", index=False)

        # light header styling on every sheet
        from openpyxl.styles import Font, PatternFill
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="0B3D66")
        for ws in writer.book.worksheets:
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
            for col in ws.columns:
                max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max(12, max_len + 2), 40)

    buf.seek(0)
    return buf.read()


def build_pdf_report(header_name: str, period_label: str, kpis: dict, target_summary: pd.DataFrame,
                      insights: list) -> bytes:
    """One-page PDF summary: header, KPI table, target vs actual table, insights."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                             leftMargin=1.8 * cm, rightMargin=1.8 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleNavy", parent=styles["Title"], textColor=NAVY, fontSize=18)
    h2_style = ParagraphStyle("H2Navy", parent=styles["Heading2"], textColor=NAVY, spaceBefore=14)
    body_style = styles["BodyText"]

    story = [
        Paragraph("Partner 360 — Summary Report", title_style),
        Paragraph(f"{header_name} &nbsp;|&nbsp; {period_label} &nbsp;|&nbsp; Generated {date.today():%d %b %Y}", body_style),
        Spacer(1, 12),
    ]

    # KPI table
    story.append(Paragraph("Key Performance Indicators", h2_style))
    kpi_rows = [
        ["Sales", format_inr(kpis.get("sales", 0))],
        ["SIP", format_inr(kpis.get("sip", 0))],
        ["AUM (net flow)", format_inr(kpis.get("aum", 0))],
        ["Revenue", format_inr(kpis.get("revenue", 0))],
        ["Total Clients", format_count(kpis.get("total_clients", 0))],
        ["Active Clients", format_count(kpis.get("active_clients", 0))],
        ["New Clients", format_count(kpis.get("new_clients", 0))],
    ]
    kpi_table = Table([["Metric", "Value"]] + kpi_rows, colWidths=[7 * cm, 7 * cm])
    kpi_table.setStyle(_table_style())
    story.append(kpi_table)

    # Target vs Actual table
    story.append(Paragraph("Target vs Achievement", h2_style))
    tva_rows = [["Metric", "Target", "Actual", "Achievement %", "Gap"]]
    for _, r in target_summary.iterrows():
        is_count = r["Metric"] == "New Clients"
        fmt = format_count if is_count else format_inr
        tva_rows.append([
            r["Metric"], fmt(r["Target"]) if pd.notna(r["Target"]) else "—",
            fmt(r["Actual"]), format_pct(r["Achievement %"]) if pd.notna(r["Achievement %"]) else "—",
            fmt(r["Gap"]) if pd.notna(r["Gap"]) else "—",
        ])
    tva_table = Table(tva_rows, colWidths=[3.5 * cm, 3 * cm, 3 * cm, 3 * cm, 3 * cm])
    tva_table.setStyle(_table_style())
    story.append(tva_table)

    # Insights
    story.append(Paragraph("Business Insights", h2_style))
    for tag, text in insights:
        story.append(Paragraph(f"<b>{tag}:</b> {text}", body_style))
        story.append(Spacer(1, 4))

    doc.build(story)
    buf.seek(0)
    return buf.read()


def _table_style():
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ])
